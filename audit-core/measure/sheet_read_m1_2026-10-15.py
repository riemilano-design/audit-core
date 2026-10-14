# -*- coding: utf-8 -*-
"""新領域① M-1: 実物の決算 Excel を**標準ライブラリだけ**で読む(封 §4)。

出すのは**読み取りの数だけ**。接続も合否も扱わない。**精度は主張しない。**
人間の決定(2026-10-15): 読み取りの依存は案 (A)(標準ライブラリ)。
**式のセルの計算結果が無ければ「計算結果なし」のまま扱い、独自計算で埋めない。**

    read   3 ファイルを読み、fixtures/sheet/m1_read_2026-10-15.json に数を書く

材料は `~/.cache/audit-core/sheet/`(封 docs/DOMAIN1_FETCH_SEAL_2026-10-15.md)。通信しない。
出典: 「令和６年度東京都各会計歳入歳出決算書及び東京都各会計歳入歳出決算附属書類等」、
東京都会計管理局、CC BY 4.0 を機械で読み取って加工。
"""
import collections
import hashlib
import io
import json
import os
import re
import sys
import zipfile
import xml.etree.ElementTree as ET

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
OPENED = os.path.join(ROOT, "fixtures", "sheet", "opened_2026-10-15.json")
OUT = os.path.join(ROOT, "fixtures", "sheet", "m1_read_2026-10-15.json")
CACHE = os.path.expanduser("~/.cache/audit-core/sheet")

NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
RELNS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
PKGREL = "{http://schemas.openxmlformats.org/package/2006/relationships}"

# 組み込みの表示形式(下見で 1 件だけ解けなかったもの。ECMA-376 の固定表)
BUILTIN = {
    0: "General", 1: "0", 2: "0.00", 3: "#,##0", 4: "#,##0.00", 9: "0%", 10: "0.00%",
    11: "0.00E+00", 12: "# ?/?", 13: "# ??/??", 14: "mm-dd-yy", 15: "d-mmm-yy",
    16: "d-mmm", 17: "mmm-yy", 18: "h:mm AM/PM", 19: "h:mm:ss AM/PM", 20: "h:mm",
    21: "h:mm:ss", 22: "m/d/yy h:mm", 37: "#,##0 ;(#,##0)", 38: "#,##0 ;[Red](#,##0)",
    39: "#,##0.00;(#,##0.00)", 40: "#,##0.00;[Red](#,##0.00)", 45: "mm:ss",
    46: "[h]:mm:ss", 47: "mmss.0", 48: "##0.0E+0", 49: "@",
}
DATEISH = re.compile(r"(?<!\\)[ymdhs]")


def _text(el):
    return "".join(el.itertext()) if el is not None else None


def _si_text(si):
    """共有文字列 1 件の**本文だけ**を取る。

    `<rPh>`(ふりがな)は本文ではない。`itertext()` を使うと
    「会計別カイケイベツ」のように**ルビが値に連結される**(M-1 で実際に踏んだ)。
    返すのは `<si><t>` と `<si><r><t>` の連結だけ。
    """
    parts = []
    for child in si:
        if child.tag == NS + "t":
            parts.append(child.text or "")
        elif child.tag == NS + "r":
            for t in child.findall(NS + "t"):
                parts.append(t.text or "")
        # rPh(ふりがな)・phoneticPr は本文ではないので取らない
    return "".join(parts)


def read(path):
    """1 冊を読む。返すのはシートごとの cells と、読み取りで気づいたことの印。"""
    z = zipfile.ZipFile(path)
    names = set(z.namelist())
    notes = {"has_shared_strings": "xl/sharedStrings.xml" in names,
             "inline_strings": 0, "shared_formulas": 0, "date1904": False,
             "unresolved_number_format_ids": 0, "errors": []}

    shared = []
    if notes["has_shared_strings"]:
        sis = ET.fromstring(z.read("xl/sharedStrings.xml")).findall(NS + "si")
        shared = [_si_text(si) for si in sis]
        notes["phonetic_strings"] = sum(1 for si in sis if si.find(NS + "rPh") is not None)
        notes["shared_strings"] = len(sis)

    fmts = {}
    if "xl/styles.xml" in names:
        st = ET.fromstring(z.read("xl/styles.xml"))
        custom = {n.get("numFmtId"): n.get("formatCode") for n in st.iter(NS + "numFmt")}
        cellxfs = st.find(NS + "cellXfs")
        for i, xf in enumerate(cellxfs.findall(NS + "xf") if cellxfs is not None else []):
            fid = xf.get("numFmtId")
            code = custom.get(fid)
            if code is None:
                code = BUILTIN.get(int(fid)) if fid is not None and fid.isdigit() else None
            if code is None:
                notes["unresolved_number_format_ids"] += 1
                code = "numFmtId:%s" % fid
            fmts[str(i)] = code

    wb = ET.fromstring(z.read("xl/workbook.xml"))
    pr = wb.find(NS + "workbookPr")
    if pr is not None and pr.get("date1904") in ("1", "true"):
        notes["date1904"] = True
    rels = {r.get("Id"): r.get("Target") for r in
            ET.fromstring(z.read("xl/_rels/workbook.xml.rels")).findall(PKGREL + "Relationship")}

    sheets = []
    for sh in wb.find(NS + "sheets").findall(NS + "sheet"):
        t = rels[sh.get(RELNS + "id")].lstrip("/")
        sheets.append((sh.get("name"), t if t.startswith("xl/") else "xl/" + t))

    out = collections.OrderedDict()
    for name, target in sheets:
        try:
            root = ET.fromstring(z.read(target))
        except Exception as e:
            notes["errors"].append({"sheet": name, "part": target, "error": type(e).__name__})
            continue
        cells, merged = {}, [mc.get("ref") for mc in root.iter(NS + "mergeCell")]
        for c in root.iter(NS + "c"):
            ref, t, s = c.get("r"), c.get("t"), c.get("s")
            f, v, isel = c.find(NS + "f"), c.find(NS + "v"), c.find(NS + "is")
            try:
                if t == "s" and v is not None:
                    value = shared[int(v.text)]
                elif t == "inlineStr":
                    value = _text(isel)
                    notes["inline_strings"] += 1
                elif t in ("str", "e"):
                    value = v.text if v is not None else None
                elif t == "b" and v is not None:
                    value = bool(int(v.text))
                elif v is not None and v.text is not None:
                    value = float(v.text)
                    if value == int(value):
                        value = int(value)
                else:
                    value = None
            except Exception as e:
                notes["errors"].append({"sheet": name, "ref": ref, "error": type(e).__name__})
                continue
            formula = None
            if f is not None:
                if f.get("t") == "shared":
                    notes["shared_formulas"] += 1
                    formula = ("=" + f.text) if f.text else {"shared_si": f.get("si")}
                elif f.text:
                    formula = "=" + f.text
            if value is None and formula is None:
                continue
            cells[ref] = {"value": value, "formula": formula,
                          "number_format": fmts.get(s) if s is not None else None,
                          "type": t}
        out[name] = {"cells": cells, "merged": merged}
    return out, notes


def summarize(label, path, book, notes):
    sheets = []
    fmt_counter = collections.Counter()
    formula_cells = no_value_formula = dateish = 0
    for name, sh in book.items():
        cells = sh["cells"]
        nf = sum(1 for c in cells.values() if c["formula"] is not None)
        nnv = sum(1 for c in cells.values() if c["formula"] is not None and c["value"] is None)
        for c in cells.values():
            fmt_counter[c["number_format"]] += 1
            if c["number_format"] and DATEISH.search(c["number_format"]):
                dateish += 1
        formula_cells += nf
        no_value_formula += nnv
        sheets.append({"name": name, "cells": len(cells), "formula_cells": nf,
                       "formula_cells_without_value": nnv, "merged": len(sh["merged"])})
    return {"label": label, "file": os.path.basename(path),
            "sha256": hashlib.sha256(io.open(path, "rb").read()).hexdigest(),
            "sheets": len(book), "cells": sum(s["cells"] for s in sheets),
            "formula_cells": formula_cells,
            "formula_cells_without_value": no_value_formula,
            "date_formatted_cells": dateish,
            "number_formats": dict(sorted(((str(k), v) for k, v in fmt_counter.items()),
                                          key=lambda kv: (-kv[1], kv[0]))),
            "notes": notes, "per_sheet": sheets}


def preview(book, sheet_index=0, rows=12, cols=8):
    """人が材料を見るための抜粋(値だけ。**判定には使わない**)。"""
    name = list(book)[sheet_index]
    cells = book[name]["cells"]
    grid = []
    for r in range(1, rows + 1):
        line = []
        for ci in range(cols):
            ref = "%s%d" % (chr(ord("A") + ci), r)
            c = cells.get(ref)
            line.append(None if c is None else (c["formula"] or c["value"]))
        grid.append(line)
    return {"sheet": name, "grid": grid}


def cross_check(path, book):
    """**実物で** openpyxl と 1 セルずつ突き合わせる(下見で「実物は未測定」と書いた所)。

    openpyxl を採用するためではない(依存は案 A で決定済み)。
    **標準ライブラリの読み取りが実物で何セル一致するか**を測るためだけに使う。
    openpyxl は式と値を同時に返せないので 2 回読む。
    """
    from openpyxl import load_workbook
    wbf = load_workbook(path, data_only=False, read_only=True)
    wbv = load_workbook(path, data_only=True, read_only=True)
    same = dv = dfmt = 0
    rows = []
    for name in wbf.sheetnames:
        cells = book.get(name, {}).get("cells", {})
        wf, wv = wbf[name], wbv[name]
        vals = {}
        for row in wv.iter_rows():
            for c in row:
                if c.value is not None:
                    vals[c.coordinate] = c.value
        fmts = {}
        for row in wf.iter_rows():
            for c in row:
                if c.value is not None:
                    fmts[c.coordinate] = c.number_format
        for ref in set(cells) | set(vals):
            g = cells.get(ref, {"value": None, "number_format": None})
            o_v = vals.get(ref)
            o_fmt = fmts.get(ref)
            ok_v = g["value"] == o_v
            ok_fmt = (g["number_format"] == o_fmt) or (o_fmt is None) or (
                o_fmt == "General" and g["number_format"] in (None, "General"))
            if ok_v and ok_fmt:
                same += 1
                continue
            dv += 0 if ok_v else 1
            dfmt += 0 if ok_fmt else 1
            if len(rows) < 20:
                rows.append({"sheet": name, "ref": ref, "std": g["value"], "openpyxl": o_v,
                             "std_format": g["number_format"], "openpyxl_format": o_fmt})
    wbf.close()
    wbv.close()
    return {"compared": same + dv + dfmt, "same": same, "value_diff": dv,
            "format_diff": dfmt, "examples": rows}


def main():
    opened = json.load(io.open(OPENED, encoding="utf-8"))
    files = opened["files"]
    books, out = {}, []
    for name, meta in files.items():
        path = os.path.join(CACHE, name)
        if not os.path.exists(path):
            raise SystemExit("材料が無い: %s(取得はしない)" % name)
        book, notes = read(path)
        books[name] = book
        s = summarize(meta["label"], path, book, notes)
        if s["sha256"] != meta["sha256"]:
            raise AssertionError("%s: sha256 が取得の記録と違う" % name)
        s["cross_check_openpyxl"] = cross_check(path, book)
        out.append(s)
        print("[%s] シート %d / セル %d / 式 %d(うち計算結果なし %d)/ 結合 %d / 共有数式 %d / 読み取り失敗 %d"
              % (s["label"], s["sheets"], s["cells"], s["formula_cells"],
                 s["formula_cells_without_value"], sum(x["merged"] for x in s["per_sheet"]),
                 notes["shared_formulas"], len(notes["errors"])))
        cc = s["cross_check_openpyxl"]
        print("    openpyxl と突き合わせ: 比べた %d / 一致 %d / 値の差 %d / 表示形式の差 %d"
              % (cc["compared"], cc["same"], cc["value_diff"], cc["format_diff"]))
    rec = {"purpose": "M-1 読み取りの結果。**精度の主張ではない**(封 §4)",
           "seal": "docs/DOMAIN1_FETCH_SEAL_2026-10-15.md",
           "source": opened["source"], "license": opened["license"],
           "dataset": opened["dataset"], "reader": "標準ライブラリのみ(案 A)",
           "files": out,
           "previews": [preview(books[n]) for n in files]}
    with io.open(OUT, "w", encoding="utf-8") as fh:
        json.dump(rec, fh, ensure_ascii=False, indent=1)
    print("->", OUT)


if __name__ == "__main__":
    main()
