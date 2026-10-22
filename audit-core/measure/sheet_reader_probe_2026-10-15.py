# -*- coding: utf-8 -*-
"""新領域① の下見: xlsx を**標準ライブラリだけ**で読めるかを測る(2026-10-15)。

設計案 docs/DOMAIN1_SHEET_DESIGN_2026-10-15.md §9 の案 (A) の費用を、意見ではなく数で出す。
**材料は自分で作る合成ブック 2 冊**(実物ではない。実物の書式は測っていない):
  1 冊目 openpyxl が書いたブック(結合セル・SUM の式・% の表示形式・inline の文字列)
  2 冊目 手で組んだ最小のブック(**実物の Excel が使う sharedStrings の経路**と、
          式と計算結果の両方を持つセル)

出すのは「標準ライブラリの読み取りが openpyxl と何セル一致したか」だけ。
**精度の主張ではない。**  usage: python3 measure/sheet_reader_probe_2026-10-15.py
"""
import os
import tempfile
import re
import zipfile
import xml.etree.ElementTree as ET

NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
RELNS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
PKGREL = "{http://schemas.openxmlformats.org/package/2006/relationships}"


def _text(el):
    return "".join(el.itertext()) if el is not None else None


def read(path):
    z = zipfile.ZipFile(path)
    names = set(z.namelist())

    shared = []
    if "xl/sharedStrings.xml" in names:
        root = ET.fromstring(z.read("xl/sharedStrings.xml"))
        shared = [_text(si) for si in root.findall(NS + "si")]

    fmts = {}
    if "xl/styles.xml" in names:
        st = ET.fromstring(z.read("xl/styles.xml"))
        custom = {n.get("numFmtId"): n.get("formatCode")
                  for n in st.iter(NS + "numFmt")}
        cellxfs = st.find(NS + "cellXfs")
        xfs = list(cellxfs.findall(NS + "xf")) if cellxfs is not None else []
        for i, xf in enumerate(xfs):
            fid = xf.get("numFmtId")
            fmts[str(i)] = custom.get(fid, fid)

    wb = ET.fromstring(z.read("xl/workbook.xml"))
    rels = {r.get("Id"): r.get("Target")
            for r in ET.fromstring(z.read("xl/_rels/workbook.xml.rels")).findall(PKGREL + "Relationship")}
    sheets = []
    for sh in wb.find(NS + "sheets").findall(NS + "sheet"):
        target = rels[sh.get(RELNS + "id")]
        target = target.lstrip("/")
        if not target.startswith("xl/"):
            target = "xl/" + target
        sheets.append((sh.get("name"), target))

    out = {}
    for name, target in sheets:
        root = ET.fromstring(z.read(target))
        cells, merged = {}, []
        for mc in root.iter(NS + "mergeCell"):
            merged.append(mc.get("ref"))
        for c in root.iter(NS + "c"):
            ref, t, s = c.get("r"), c.get("t"), c.get("s")
            f = c.find(NS + "f")
            v = c.find(NS + "v")
            isel = c.find(NS + "is")
            if t == "s" and v is not None:
                value = shared[int(v.text)]
            elif t == "inlineStr":
                value = _text(isel)
            elif t == "str":
                value = v.text if v is not None else None
            elif t == "b" and v is not None:
                value = bool(int(v.text))
            elif v is not None and v.text is not None:
                value = float(v.text)
                if value == int(value):
                    value = int(value)
            else:
                value = None
            if value is None and f is None:
                continue
            cells[ref] = {"value": value,
                          "formula": ("=" + f.text) if (f is not None and f.text) else None,
                          "number_format": fmts.get(s) if s is not None else None}
        out[name] = {"cells": cells, "merged": merged}
    return out




def make_openpyxl_book(path):
    """openpyxl が書いたブック(結合セル・SUM の式・% の表示形式)。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment
    wb = Workbook()
    ws = wb.active
    ws.title = "月次"
    ws["A1"] = "部門別売上(単位:千円)"
    ws.merge_cells("A1:D1")
    ws["A1"].alignment = Alignment(horizontal="left")
    for ref, val in (("A2", "項目"), ("B2", "当月"), ("C2", "累計"), ("D2", "前年同月")):
        ws[ref] = val
    rows = [("東京支店", 12000, 48000, 11000), ("大阪支店", 8000, 31000, 8500),
            ("名古屋支店", 5000, 19000, 4800)]
    for i, (name, a, b, c) in enumerate(rows, start=3):
        ws["A%d" % i], ws["B%d" % i], ws["C%d" % i], ws["D%d" % i] = name, a, b, c
    ws["A6"] = "小計(国内)"
    for col in "BCD":
        ws["%s6" % col] = "=SUM(%s3:%s5)" % (col, col)
    ws["A7"] = "海外"
    ws["B7"], ws["C7"], ws["D7"] = 3000, 9000, 2500
    ws["A8"] = "合計"
    for col in "BCD":
        ws["%s8" % col] = "=%s6+%s7" % (col, col)
    ws["B9"] = 0.0825
    ws["B9"].number_format = "0.00%"
    ws2 = wb.create_sheet("年次")
    ws2["A1"] = "単位:百万円"
    ws2["A2"], ws2["B2"], ws2["C2"] = "項目", "当期", "前期"
    ws2["A3"], ws2["B3"], ws2["C3"] = "売上高", 68, 61
    ws2["A4"], ws2["B4"], ws2["C4"] = "売上総利益", 12, 11
    wb.save(path)


_S = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_P = "http://schemas.openxmlformats.org/package/2006/relationships"


def make_shared_strings_book(path):
    """実物の Excel が使う **sharedStrings** の経路と、式 + 計算結果を持つセル。"""
    parts = {
        "[Content_Types].xml":
            '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
            '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
            '<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>',
        "_rels/.rels":
            '<?xml version="1.0"?><Relationships xmlns="%s"><Relationship Id="rId1" '
            'Type="%s/officeDocument" Target="xl/workbook.xml"/></Relationships>' % (_P, _R.rsplit("/", 1)[0]),
        "xl/workbook.xml":
            '<?xml version="1.0"?><workbook xmlns="%s" xmlns:r="%s"><sheets>'
            '<sheet name="決算" sheetId="1" r:id="rId1"/></sheets></workbook>' % (_S, _R),
        "xl/_rels/workbook.xml.rels":
            '<?xml version="1.0"?><Relationships xmlns="%s">'
            '<Relationship Id="rId1" Type="%s/worksheet" Target="worksheets/sheet1.xml"/>'
            '<Relationship Id="rId2" Type="%s/sharedStrings" Target="sharedStrings.xml"/></Relationships>' % (_P, _R, _R),
        "xl/sharedStrings.xml":
            '<?xml version="1.0"?><sst xmlns="%s" count="2" uniqueCount="2">'
            '<si><t>売上高</t></si><si><t>単位:千円</t></si></sst>' % _S,
        "xl/worksheets/sheet1.xml":
            '<?xml version="1.0"?><worksheet xmlns="%s"><sheetData>'
            '<row r="1"><c r="A1" t="s"><v>1</v></c></row>'
            '<row r="2"><c r="A2" t="s"><v>0</v></c><c r="B2"><v>9648</v></c></row>'
            '<row r="3"><c r="B3"><f>SUM(B2:B2)</f><v>9648</v></c></row>'
            '</sheetData></worksheet>' % _S,
    }
    with zipfile.ZipFile(path, "w") as z:
        for name, body in parts.items():
            z.writestr(name, body)


def compare(path):
    """標準ライブラリの読み取りと openpyxl を 1 セルずつ比べる。

    openpyxl は**式と値を同時に返せない**ので 2 回読む(`data_only` の True/False)。
    返すのは (一致したセル数, 値の差, 式の差, 表示形式の差, 差の明細)。
    """
    from openpyxl import load_workbook
    std = read(path)
    wbf = load_workbook(path, data_only=False)
    wbv = load_workbook(path, data_only=True)
    same = dv = df = dfmt = 0
    rows = []
    for sh in wbf.sheetnames:
        wf, wv, cells = wbf[sh], wbv[sh], std[sh]["cells"]
        refs = set(cells)
        for row in wf.iter_rows():
            for c in row:
                if c.value is not None:
                    refs.add(c.coordinate)
        for ref in sorted(refs):
            o_f, o_v, o_fmt = wf[ref].value, wv[ref].value, wf[ref].number_format
            g = cells.get(ref, {"value": None, "formula": None, "number_format": None})
            exp_f = o_f if (isinstance(o_f, str) and o_f.startswith("=")) else None
            ok_v = (g["value"] == o_v) or (exp_f is None and g["value"] == o_f)
            ok_f = (g["formula"] == exp_f)
            ok_fmt = (g["number_format"] == o_fmt) or (o_fmt == "General" and g["number_format"] in (None, "0"))
            if ok_v and ok_f and ok_fmt:
                same += 1
                continue
            dv += 0 if ok_v else 1
            df += 0 if ok_f else 1
            dfmt += 0 if ok_fmt else 1
            rows.append({"sheet": sh, "ref": ref, "std": g,
                         "openpyxl": {"value": o_v, "formula": exp_f, "number_format": o_fmt}})
    return {"same": same, "value_diff": dv, "formula_diff": df, "format_diff": dfmt, "rows": rows}


def main():
    d = tempfile.mkdtemp(prefix="sheet-probe-")
    a = os.path.join(d, "openpyxl_written.xlsx")
    b = os.path.join(d, "shared_strings.xlsx")
    make_openpyxl_book(a)
    make_shared_strings_book(b)
    for name, path in (("openpyxl が書いたブック", a), ("sharedStrings のブック", b)):
        r = compare(path)
        print("[%s] 一致 %d セル / 値の差 %d / 式の差 %d / 表示形式の差 %d"
              % (name, r["same"], r["value_diff"], r["formula_diff"], r["format_diff"]))
        for row in r["rows"]:
            print("    差: %s!%s std=%r openpyxl=%r"
                  % (row["sheet"], row["ref"], row["std"], row["openpyxl"]))
    # 式のセルの計算結果は、保存したアプリが入れていなければ**無い**
    from openpyxl import load_workbook
    wv = load_workbook(a, data_only=True)["月次"]
    print("openpyxl data_only=True が式のセルで返した値(openpyxl が保存した本):",
          {c: wv[c].value for c in ("B6", "B8")})


if __name__ == "__main__":
    main()
